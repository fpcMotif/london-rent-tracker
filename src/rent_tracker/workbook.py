import json
import sqlite3
from dataclasses import astuple, fields
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from rent_tracker.config import Config
from rent_tracker.store import coverage_rows, listing_rows, run_dates, run_listings
from rent_tracker.summary import AreaSummary, summarise

LISTING_COLUMNS = [
    "run_date",
    "source",
    "listing_id",
    "area",
    "address",
    "postcode_district",
    "price_pcm",
    "beds",
    "baths",
    "sq_ft",
    "status",
    "first_seen",
    "last_seen",
    "url",
    "canonical",
    "other_source_ids",
    "transit_min",
    "within_commute",
    "cycle_min",
    "route",
]


def within_label(transit_min: int | None, limit: int) -> str:
    if transit_min is None:
        return ""
    return "yes" if transit_min <= limit else "no"


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for col, name in zip(ws.columns, header):
        width = max(
            [len(str(name))] + [len(str(c.value or "")) for c in list(col)[:200]]
        )
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 60)


def summaries(db: sqlite3.Connection, cfg: Config) -> list[AreaSummary]:
    out = []
    for run_date in run_dates(db):
        rows = run_listings(db, run_date)
        for area in cfg.areas:
            out.append(
                summarise(
                    run_date,
                    area.name,
                    [r for r in rows if r["area"] == area.name],
                    cfg.destination.max_commute_min,
                )
            )
    return out


def build(db: sqlite3.Connection, cfg: Config, path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    listing_values = []
    for r in listing_rows(db):
        row = dict(zip(r.keys(), r, strict=True))
        row["canonical"] = "yes" if row["canonical"] else "no"
        row["other_source_ids"] = ", ".join(json.loads(row["other_source_ids"]))
        row["within_commute"] = within_label(
            row["transit_min"], cfg.destination.max_commute_min
        )
        listing_values.append([row[c] for c in LISTING_COLUMNS])
    _sheet(wb, "listings", LISTING_COLUMNS, listing_values)

    summary_header = [f.name for f in fields(AreaSummary)]
    _sheet(
        wb,
        "weekly_summary",
        summary_header,
        [list(astuple(s)) for s in summaries(db, cfg)],
    )

    cov = coverage_rows(db)
    _sheet(
        wb,
        "coverage",
        [
            "run_date",
            "source",
            "area",
            "search",
            "displayed",
            "collected",
            "status",
            "note",
        ],
        [
            [
                r["run_date"],
                r["source"],
                r["area"],
                r["search"],
                r["displayed"],
                r["collected"],
                r["status"],
                r["note"],
            ]
            for r in cov
        ],
    )

    commute = db.execute("SELECT * FROM commute_points ORDER BY transit_min").fetchall()
    _sheet(
        wb,
        "commute",
        [
            "lat",
            "lng",
            "transit_min",
            "within_commute",
            "cycle_min",
            "route",
            "checked_on",
        ],
        [
            [
                r["lat"],
                r["lng"],
                r["transit_min"],
                within_label(r["transit_min"], cfg.destination.max_commute_min),
                r["cycle_min"],
                r["route"],
                r["checked_on"],
            ]
            for r in commute
        ],
    )

    ons = db.execute("SELECT * FROM ons ORDER BY borough, month").fetchall()
    _sheet(
        wb,
        "ons_benchmark",
        ["month", "borough", "two_bed_pcm", "two_bed_annual_change_pct", "all_pcm"],
        [
            [
                r["month"],
                r["borough"],
                r["two_bed_pcm"],
                r["two_bed_annual_change"],
                r["all_pcm"],
            ]
            for r in ons
        ],
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)

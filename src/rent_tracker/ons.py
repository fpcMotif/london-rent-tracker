import io
import re
import sqlite3

import httpx
from openpyxl import load_workbook

from rent_tracker.config import Config

DATASET = "https://www.ons.gov.uk/economy/inflationandpriceindices/datasets/priceindexofprivaterentsukmonthlypricestatistics"


def latest_file_url(client: httpx.Client) -> str:
    page = client.get(DATASET)
    page.raise_for_status()
    m = re.search(r'href="(/file\?uri=[^"]+\.xlsx)"', page.text)
    if not m:
        raise RuntimeError("no PIPR xlsx link on the ONS dataset page")
    return "https://www.ons.gov.uk" + m.group(1)


def _number(v: object) -> float | None:
    return float(v) if isinstance(v, (int, float)) else None


def refresh(
    db: sqlite3.Connection, client: httpx.Client, cfg: Config, months: int = 13
) -> str:
    url = latest_file_url(client)
    resp = client.get(url)
    resp.raise_for_status()
    ws = load_workbook(io.BytesIO(resp.content), read_only=True)["Table 1"]
    rows = ws.iter_rows(values_only=True)
    header = None
    for row in rows:
        if row and row[0] == "Time period":
            header = list(row)
            break
    if header is None:
        raise RuntimeError("PIPR Table 1 header not found")
    col = {name: i for i, name in enumerate(header)}
    boroughs = {b for a in cfg.areas for b in a.ons_boroughs}
    found = []
    for row in rows:
        if row[col["Area name"]] in boroughs:
            found.append(row)
    keep_months = sorted({r[col["Time period"]] for r in found})[-months:]
    with db:
        for r in found:
            if r[col["Time period"]] not in keep_months:
                continue
            two_bed = _number(r[col["Rental price two bed"]])
            change = _number(r[col["Annual change two bed"]])
            everything = _number(r[col["Rental price"]])
            db.execute(
                "INSERT OR REPLACE INTO ons VALUES (?,?,?,?,?)",
                (
                    r[col["Time period"]].strftime("%Y-%m"),
                    r[col["Area name"]],
                    round(two_bed) if two_bed is not None else None,
                    round(change, 2) if change is not None else None,
                    round(everything) if everything is not None else None,
                ),
            )
    return url
